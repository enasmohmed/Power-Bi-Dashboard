"""
Parse the "Daily Tracker" sheet from a workbook (xlsx / xlsm).

Expected layout (flexible):
- A sheet named exactly "Daily Tracker" (case-insensitive).
- A header row containing "Picker Name" (case-insensitive).
- Columns after the picker column: per-day line counts (numbers, dashes, or empty).
- Date headers like "01-Oct" with year taken from a title row above (e.g. "October 2025").
"""

from __future__ import annotations

import io
import re
from calendar import month_name
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

GREEN_MIN = 200
YELLOW_MIN = 50

_MONTH_YEAR = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)"
    r"\s*[,|\s]+\s*(20\d{2})",
    re.I,
)
_YEAR_ANY = re.compile(r"\b(20\d{2})\b")


def _norm(cell: Any) -> str:
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return str(cell)
    if isinstance(cell, float) and cell == int(cell):
        return str(int(cell))
    return str(cell).strip()


def _parse_cell_number(cell: Any) -> int | None:
    if cell is None or cell == "":
        return None
    if isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        if isinstance(cell, float):
            return int(round(cell))
        return int(cell)
    s = str(cell).strip()
    if s in ("-", "—", "–", "N/A", "n/a"):
        return None
    try:
        return int(float(s.replace(",", "")))
    except ValueError:
        return None


def _infer_year_from_rows(rows_before: list[tuple[Any, ...]]) -> int:
    blob = " ".join(
        " ".join(_norm(c) for c in row if _norm(c)) for row in rows_before if row
    )
    m = _MONTH_YEAR.search(blob)
    if m:
        return int(m.group(2))
    m2 = _YEAR_ANY.search(blob)
    if m2:
        return int(m2.group(1))
    return datetime.now().year


def _extract_title_line(rows_before: list[tuple[Any, ...]]) -> str:
    parts: list[str] = []
    for row in rows_before:
        if not row:
            continue
        line = " ".join(_norm(c) for c in row if _norm(c))
        if line:
            parts.append(line)
    return " | ".join(parts[:4]) if parts else "Daily lines tracker"


def _parse_date_header(cell: Any, default_year: int) -> date | None:
    if cell is None or cell == "":
        return None
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    s = _norm(cell)
    if s in ("-", "—"):
        return None
    for fmt in ("%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    for fmt in ("%d-%b", "%d-%b."):
        try:
            base = s.rstrip(".")
            return datetime.strptime(f"{base}-{default_year}", "%d-%b-%Y").date()
        except ValueError:
            continue
    return None


def _month_key(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def _month_label(key: str) -> str:
    try:
        y, m = key.split("-")
        mi = int(m)
        yi = int(y)
        return f"{month_name[mi]} {yi}"
    except (ValueError, IndexError):
        return key


def _tier(n: int | None) -> str:
    if n is None:
        return "empty"
    if n >= GREEN_MIN:
        return "green"
    if n >= YELLOW_MIN:
        return "yellow"
    return "red"


def tier_for_value(lines: int | None) -> str:
    """Same banding as the spreadsheet (for API edits after upload)."""
    return _tier(lines)


def parse_daily_tracker_workbook(raw_bytes: bytes) -> dict[str, Any]:
    bio = io.BytesIO(raw_bytes)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        sheet_title = None
        for name in wb.sheetnames:
            if name.strip().lower() == "daily tracker":
                sheet_title = name
                break
        if not sheet_title:
            raise ValueError('No sheet named "Daily Tracker" was found in the workbook.')

        ws = wb[sheet_title]
        data_rows: list[tuple[Any, ...]] = [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    if not data_rows:
        raise ValueError('The "Daily Tracker" sheet is empty.')

    header_idx: int | None = None
    picker_col: int | None = None
    search_limit = min(len(data_rows), 300)
    for i in range(search_limit):
        row = data_rows[i]
        if not row:
            continue
        for j, cell in enumerate(row):
            if _norm(cell).lower() == "picker name":
                header_idx = i
                picker_col = j
                break
        if header_idx is not None:
            break

    if header_idx is None or picker_col is None:
        raise ValueError(
            'Could not find a "Picker Name" column header on the "Daily Tracker" sheet.'
        )

    rows_before = list(data_rows[:header_idx])
    default_year = _infer_year_from_rows(rows_before)
    title_line = _extract_title_line(rows_before)

    header_row = list(data_rows[header_idx])
    date_columns: list[tuple[int, date, str]] = []
    for j in range(picker_col + 1, len(header_row)):
        raw_h = header_row[j]
        d = _parse_date_header(raw_h, default_year)
        if d is None:
            continue
        label = d.strftime("%d-%b")
        date_columns.append((j, d, label))

    if not date_columns:
        raise ValueError(
            "No date columns were found after Picker Name. Expected headers like 01-Oct with a year in the title area."
        )

    month_keys: list[str] = []
    seen_m: set[str] = set()
    for _, d, _ in date_columns:
        mk = _month_key(d)
        if mk not in seen_m:
            seen_m.add(mk)
            month_keys.append(mk)
    month_keys.sort()
    months = [{"key": mk, "label": _month_label(mk)} for mk in month_keys]

    columns_out: list[dict[str, Any]] = []
    for j, d, label in date_columns:
        columns_out.append(
            {
                "iso": d.isoformat(),
                "label": label,
                "month": _month_key(d),
            }
        )

    rows_out: list[dict[str, Any]] = []
    pickers_set: set[str] = set()
    for row in data_rows[header_idx + 1 :]:
        if not row or picker_col >= len(row):
            continue
        pname = _norm(row[picker_col])
        if not pname or pname.lower() in ("total", "totals", "sum", "average", "avg"):
            continue
        values: list[int | None] = []
        tiers: list[str] = []
        for j, _d, _lbl in date_columns:
            cell = row[j] if j < len(row) else None
            n = _parse_cell_number(cell)
            values.append(n)
            tiers.append(_tier(n))
        if all(v is None for v in values):
            continue
        pickers_set.add(pname)
        rows_out.append({"picker": pname, "values": values, "tiers": tiers})

    if not rows_out:
        raise ValueError("No data rows with line counts were found under the header.")

    pickers_sorted = sorted(pickers_set, key=lambda s: s.lower())

    return {
        "table_title": title_line.upper() if title_line else "DAILY LINES TRACKER",
        "legend": {
            "green_ge": GREEN_MIN,
            "yellow_ge": YELLOW_MIN,
            "red_lt": YELLOW_MIN,
        },
        "thresholds": {"green": GREEN_MIN, "yellow_min": YELLOW_MIN},
        "months": months,
        "pickers": pickers_sorted,
        "columns": columns_out,
        "rows": rows_out,
    }
